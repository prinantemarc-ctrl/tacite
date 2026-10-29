"""Production de fichiers (CSV, XLSX, Markdown, DOCX, HTML, PDF) et stockage des artefacts.

Le rendu HTML/PDF reprend le design Gorgone : papier, encre, bronze, filets, étiquettes en capitales.
"""

from __future__ import annotations

import base64
import csv
import io
import logging
import math
import re
import uuid
from datetime import datetime
from pathlib import Path

from sqlalchemy.orm import Session

from ..config import ARTIFACTS_DIR
from ..constants import AXES
from ..i18n import axis_labels, axis_poles_map, get_locale
from ..models import Artifact

log = logging.getLogger("tacite.files")

MIME = {
    "csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "md": "text/markdown; charset=utf-8", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "html": "text/html; charset=utf-8", "pdf": "application/pdf", "png": "image/png", "svg": "image/svg+xml", "json": "application/json",
}


def _safe_name(name: str, ext: str) -> str:
    base = re.sub(r"[^\w\-. ]+", "", (name or "fichier").strip(), flags=re.U).strip().replace(" ", "_")[:80] or "fichier"
    if not base.lower().endswith("." + ext):
        base += "." + ext
    return base


def save_artifact(db: Session, kind: str, filename: str, data: bytes, title: str | None = None,
                  conversation_id: int | None = None, message_id: int | None = None, meta: dict | None = None) -> Artifact:
    root = Path(ARTIFACTS_DIR)
    root.mkdir(parents=True, exist_ok=True)
    fname = _safe_name(filename, kind)
    path = root / f"{datetime.utcnow():%Y%m%d}-{uuid.uuid4().hex[:8]}-{fname}"
    path.write_bytes(data)
    art = Artifact(conversation_id=conversation_id, message_id=message_id, kind=kind, title=title or fname,
                   filename=fname, path=str(path), mime=MIME.get(kind, "application/octet-stream"), size=len(data), meta=meta or {})
    db.add(art)
    db.commit()
    return art


# ───────────────────────────────────────────────────────────── tableaux


def _rows_columns(rows: list[dict], columns: list[str] | None) -> list[str]:
    if columns:
        return columns
    seen: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.append(k)
    return seen


def csv_bytes(rows: list[dict], columns: list[str] | None = None) -> bytes:
    cols = _rows_columns(rows, columns)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(cols)
    for r in rows:
        w.writerow([_cell(r.get(c)) for c in cols])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _cell(v):
    if isinstance(v, float):
        return f"{v:.3f}".rstrip("0").rstrip(".")
    if isinstance(v, (list, dict)):
        import json
        return json.dumps(v, ensure_ascii=False)
    return "" if v is None else v


def xlsx_bytes(rows: list[dict], columns: list[str] | None = None, sheet: str = "Données") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    cols = _rows_columns(rows, columns)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="6E6230")
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_cell(r.get(c)) for c in cols])
    for i, c in enumerate(cols, 1):
        width = max([len(str(c))] + [len(str(_cell(r.get(c)))) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ───────────────────────────────────────────────────────────── documents


def md_to_html_fragment(md: str) -> str:
    import markdown
    return markdown.markdown(md or "", extensions=["tables", "fenced_code", "sane_lists"])


_LIST_REFUSAL = re.compile(
    r"[^.!?\n]*("
    r"no client file|does not support a journalist|journalist target list|"
    r"liste de journalistes|liste serait fausse|would be false|"
    r"scores are not saturated|scores ne sont pas satur|"
    r"aucun dossier client|pas de dossier client|client ne le permettent pas"
    r")[^.!?\n]*[.!?]?",
    re.I,
)
_JKEYS = ("name", "outlet", "articles", "accessibility", "action", "contact")


def strip_list_refusal(text: str) -> str:
    """Retire la phrase qui refuse la liste de journalistes quand les noms sont connus."""
    cleaned = _LIST_REFUSAL.sub(" ", text or "")
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    cleaned = re.sub(r" *\n", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def journalist_rows(raw) -> list[dict]:
    """Normalise influence_targets, who_covers ou une fiche média vers des lignes de tableau."""
    rows: list[dict] = []
    if not isinstance(raw, list):
        return rows
    for t in raw:
        if not isinstance(t, dict):
            continue
        name = str(t.get("name") or t.get("journalist") or "").strip()
        if not name:
            continue
        bits: list[str] = []
        if isinstance(t.get("contact"), str) and t["contact"].strip():
            bits.append(t["contact"].strip())
        contacts = t.get("contacts") or []
        if isinstance(contacts, list):
            for c in contacts[:2]:
                if isinstance(c, dict):
                    val = str(c.get("value") or "").strip()
                    if val:
                        bits.append(val)
                elif c:
                    bits.append(str(c).strip())
        articles = t.get("articles_on_topic")
        if articles is None:
            articles = t.get("articles") if t.get("articles") is not None else t.get("n")
        acc = t.get("accessibility")
        if acc is None:
            acc = t.get("accessibility_score")
        seen: list[str] = []
        for b in bits:
            if b not in seen:
                seen.append(b)
        rows.append({
            "name": name,
            "outlet": t.get("outlet") or "",
            "articles": "" if articles is None else articles,
            "accessibility": "" if acc is None else acc,
            "action": t.get("recommended_action") or t.get("action") or "",
            "contact": ", ".join(seen),
            "basis": str(t.get("basis") or "").strip(),
        })
        if len(rows) >= 12:
            break
    return rows


def _jcell(v) -> str:
    if v is None or v == "":
        return "—"
    return str(v)


def _jheads(en: bool) -> list[str]:
    if en:
        return ["Journalist", "Outlet", "Articles", "Access", "Action", "Contact"]
    return ["Journaliste", "Média", "Articles", "Accès", "Action", "Contact"]


def _shared_basis(rows: list[dict]) -> str:
    bases = {(r.get("basis") or "").strip() for r in rows}
    bases.discard("")
    if len(bases) == 1 and all((r.get("basis") or "").strip() for r in rows):
        return next(iter(bases))
    return ""


def journalist_markdown(rows: list[dict]) -> str:
    en = get_locale() == "en"
    heads = _jheads(en)
    lines = [
        "| " + " | ".join(heads) + " |",
        "| " + " | ".join("---" for _ in heads) + " |",
    ]
    for r in rows:
        cells = [_jcell(r.get(k)).replace("|", "/") for k in _JKEYS]
        lines.append("| " + " | ".join(cells) + " |")
    cap = _shared_basis(rows)
    return "\n".join(lines) + (f"\n\n{cap}" if cap else "")


def journalist_table_html(rows: list[dict], en: bool) -> str:
    heads = _jheads(en)
    th = "".join(f"<th>{_esc(h)}</th>" for h in heads)
    body = []
    for r in rows:
        tds = "".join(f"<td>{_esc(_jcell(r.get(k)))}</td>" for k in _JKEYS)
        body.append(f"<tr>{tds}</tr>")
    cap = _shared_basis(rows)
    caption = f'<p class="sub">{_esc(cap)}</p>' if cap else ""
    return f"<table><thead><tr>{th}</tr></thead><tbody>{''.join(body)}</tbody></table>{caption}"


def _note_english(*chunks: str) -> bool:
    text = " ".join(c or "" for c in chunks)
    en_hits = len(re.findall(r"\b(the|and|with|through|from|this|that|does|not|file|outlet|ministry)\b", text, re.I))
    fr_hits = len(re.findall(r"\b(les|des|une|dans|avec|cette|qui|pour|média|pas|ligne)\b", text, re.I))
    if en_hits == fr_hits:
        return get_locale() == "en"
    return en_hits > fr_hits


_FONT_PATH = Path(__file__).resolve().parent.parent / "assets" / "geist.woff2"
_KPI_TABLE = re.compile(
    r"(?:^|\n)(\|[^\n]+\|\n\|[-:| ]+\|\n\|[^\n]+\|)(?:\n|$)",
)
_KPI_SPLIT = re.compile(
    r"^(\d{1,3}(?:[ \u00a0]\d{3})+|\d{1,2}\s*[–—-]\s*\d{1,2}|\d+(?:[.,]\d+)?)\s*(.*)$"
)


def _font_css() -> str:
    if not _FONT_PATH.exists():
        return ""
    b64 = base64.b64encode(_FONT_PATH.read_bytes()).decode("ascii")
    return (
        '@font-face{font-family:"Geist";src:url(data:font/woff2;base64,'
        + b64
        + ') format("woff2");font-weight:100 900;font-style:normal;font-display:swap;}'
    )


def _stat(label: str, value: str) -> str:
    """Le grand caractère est un chiffre. Le reste de la cellule passe en légende."""
    m = _KPI_SPLIT.match(value.strip())
    if m and m.group(2).strip():
        val, sub = m.group(1).strip(), m.group(2).strip()
        return (
            f'<div class="stat"><div class="lbl">{_esc(label)}</div>'
            f'<div class="val">{_esc(val)}</div><p class="sub">{_esc(sub)}</p></div>'
        )
    klass = "val long" if len(value.strip()) > 16 else "val"
    return f'<div class="stat"><div class="lbl">{_esc(label)}</div><div class="{klass}">{_esc(value)}</div></div>'


def _extract_kpi_table(md: str) -> tuple[str, str]:
    """Une table d'une seule ligne de valeurs devient la bande de chiffres du kit."""
    m = _KPI_TABLE.search(md or "")
    if not m:
        return md, ""
    rows = [r for r in m.group(1).splitlines() if not re.match(r"^\|[-:| ]+\|$", r.strip())]
    if len(rows) != 2:
        return md, ""
    cells = lambda line: [c.strip() for c in line.strip().strip("|").split("|")]
    heads, vals = cells(rows[0]), cells(rows[1])
    if not (2 <= len(heads) <= 5) or len(heads) != len(vals):
        return md, ""
    if any(len(c) > 48 or not c for c in heads + vals):
        return md, ""
    n = len(heads)
    html = f'<div class="stats cols-{n}">' + "".join(_stat(h, v) for h, v in zip(heads, vals)) + "</div>"
    return (md[:m.start()] + "\n" + md[m.end():]).strip(), html


def _classify_headings(html: str) -> str:
    def repl(m: re.Match) -> str:
        inner = m.group(1)
        text = re.sub(r"<[^>]+>", "", inner).strip()
        kind = "label" if len(text) <= 32 else "title"
        return f'<h2 class="{kind}">{inner}</h2>'
    return re.sub(r"<h2>(.*?)</h2>", repl, html, flags=re.S)


def _lead_paragraph(html: str) -> str:
    return re.sub(r"<p>", '<p class="lead">', html, count=1)


def _plain(md: str) -> str:
    """Retire HTML et CSS : une note MCP ne peut pas apporter sa propre mise en page."""
    text = re.sub(r"<style[\s\S]*?</style>", "", md or "", flags=re.I)
    text = re.sub(r"</?[a-zA-Z][^>]*>", "", text)
    return text


def render_press_note(
    title: str,
    reading: str,
    frames: str,
    evidence: str,
    effects: str,
    lever: str,
    figures: list | None = None,
    subtitle: str = "",
    kicker: str = "Tacite · Press Intelligence",
    journalists: list | None = None,
) -> str:
    """Note en cinq blocs, toujours dans le gabarit Tacite. L'appelant ne choisit pas le design."""
    rows = journalist_rows(journalists)
    en = _note_english(reading, frames, evidence, effects, lever, subtitle)
    if rows:
        lever = strip_list_refusal(lever)
    chunks: list[str] = []
    if (reading or "").strip():
        chunks.append(_lead_paragraph(md_to_html_fragment(_plain(reading))))
    stats = []
    for fig in (figures or [])[:5]:
        if isinstance(fig, dict) and fig.get("label") and fig.get("value") is not None:
            stats.append(_stat(str(fig["label"]), str(fig["value"])))
    if stats:
        chunks.append(f'<div class="stats cols-{len(stats)}">' + "".join(stats) + "</div>")
    labels = ("Frames", "Evidence", "Effects", "Lever") if en else ("Cadrages", "Preuves", "Effets", "Levier")
    bodies = (frames, evidence, effects, lever)
    for heading, body in zip(labels, bodies):
        show_table = heading in ("Lever", "Levier") and bool(rows)
        if not (body or "").strip() and not show_table:
            continue
        chunks.append(f'<h2 class="label">{_esc(heading)}</h2>')
        if (body or "").strip():
            chunks.append(md_to_html_fragment(_plain(body)))
        if show_table:
            chunks.append(journalist_table_html(rows, en))
    if en:
        method = (
            "Event and frame are kept apart. Confidence follows the strength of the sources, not the volume. "
            "The analytical frames do not judge the substance."
        )
        method_lbl = "Method"
    else:
        method = (
            "Événement et cadrage sont tenus à part. La confiance suit la solidité des sources, pas le volume. "
            "Les cadres d'analyse ne jugent pas le fond."
        )
        method_lbl = "Méthode"
    chunks.append(f'<div class="callout"><div class="lbl">{method_lbl}</div>{method}</div>')
    return html_document(title, "\n".join(chunks), kicker=kicker, subtitle=subtitle)


def render_markdown_body(md: str) -> str:
    """Markdown d'une note MCP → fragments du kit Tacite (chiffres, chapô, rubriques)."""
    md, stats = _extract_kpi_table(_plain(md or ""))
    html = _lead_paragraph(_classify_headings(md_to_html_fragment(md)))
    if not stats:
        return html
    if "</p>" in html:
        return html.replace("</p>", "</p>" + stats, 1)
    return stats + html


GORGONE_CSS = """
@page {
  size: A4;
  margin: 16mm 16mm 18mm 16mm;
  @bottom-center { content: counter(page) " / " counter(pages); font-family: "Geist", Helvetica, sans-serif; font-size: 8pt; color: #9a958b; }
}
__FONT__
:root {
  --paper: #ffffff; --ink: #1c1b19; --muted: #5c5850; --faint: #8f8b82;
  --accent: #6e6230; --accent-strong: #4f4622; --hair: #e4e1da; --hair-soft: #eeece6;
  --surface: #f4f2ed; --tint: #f7f4ea; --alert: #8c3a2f;
}
* { box-sizing: border-box; }
body {
  font-family: "Geist", -apple-system, "Segoe UI", Helvetica, Arial, sans-serif;
  color: var(--ink); background: #fff; font-size: 10.5pt; line-height: 1.55; margin: 0;
  font-variant-ligatures: none; font-feature-settings: "liga" 0, "calt" 0; -webkit-font-smoothing: antialiased;
}
.doc { max-width: 178mm; margin: 0 auto; }
.mast { display: flex; align-items: flex-end; justify-content: space-between; gap: 18px; }
.mast-main { min-width: 0; flex: 1; }
.kicker { font-size: 8pt; letter-spacing: 0; text-transform: uppercase; color: var(--accent-strong); font-weight: 650; }
h1 { font-size: 22pt; font-weight: 700; margin: 5px 0 0; letter-spacing: -0.02em; line-height: 1.12; }
.sub { color: var(--muted); font-size: 10pt; margin: 6px 0 0; }
.meta { text-align: right; font-size: 9pt; line-height: 1.45; color: var(--muted); flex: 0 0 52mm; }
.meta div:first-child { color: var(--ink); font-weight: 650; }
.rule { border: 0; border-top: 1.5px solid var(--accent); margin: 12px 0 16px; }
h2, h2.label { page-break-after: avoid; font-size: 8.5pt; letter-spacing: 0; text-transform: uppercase; color: var(--accent-strong); font-weight: 700; border-bottom: 1px solid var(--hair); padding-bottom: 3px; margin: 18px 0 8px; }
h2.title { font-size: 13pt; font-weight: 650; letter-spacing: -0.015em; text-transform: none; color: var(--ink); border-bottom: 1px solid var(--hair); padding-bottom: 3px; margin: 18px 0 8px; }
h3 { font-size: 11.5pt; margin: 12px 0 4px; font-weight: 650; }
p { margin: 0 0 8px; }
p.lead { font-size: 12pt; line-height: 1.45; margin: 0 0 14px; }
ul, ol { margin: 0 0 10px; padding-left: 1.15em; }
li { margin-bottom: 4px; }
li::marker { color: var(--accent); }
table { width: 100%; border-collapse: collapse; font-size: 9.5pt; margin: 8px 0 12px; }
th { text-align: left; font-size: 8pt; letter-spacing: 0; text-transform: uppercase; color: var(--accent-strong); border-bottom: 1px solid var(--hair); padding: 5px 6px; font-weight: 700; background: var(--tint); }
td { border-bottom: 1px solid var(--hair-soft); padding: 5px 6px; vertical-align: top; }
.num, .val { font-variant-numeric: tabular-nums; }
blockquote, .callout { border: 0; border-left: 3px solid var(--accent); background: var(--surface); padding: 8px 12px; margin: 12px 0; font-size: 10pt; color: var(--muted); page-break-inside: avoid; }
blockquote p { margin: 0; }
blockquote p + p { margin-top: 4px; }
blockquote strong { display: block; font-size: 8pt; letter-spacing: 0; text-transform: uppercase; color: var(--accent-strong); font-weight: 700; margin-bottom: 3px; }
.callout .lbl { font-size: 8pt; letter-spacing: 0; text-transform: uppercase; color: var(--accent-strong); font-weight: 700; margin-bottom: 3px; }
.stats { display: grid; gap: 10px 16px; margin: 2px 0 16px; page-break-inside: avoid; }
.stats.cols-2 { grid-template-columns: repeat(2, 1fr); }
.stats.cols-3 { grid-template-columns: repeat(3, 1fr); }
.stats.cols-4, .stats.cols-5 { grid-template-columns: repeat(4, 1fr); }
.stat { border-left: 1.5px solid var(--accent); padding: 1px 0 1px 9px; }
.stat .lbl { font-size: 7.5pt; letter-spacing: 0; text-transform: uppercase; color: var(--accent-strong); font-weight: 700; }
.stat .val { font-size: 15pt; font-weight: 700; letter-spacing: -0.02em; line-height: 1.15; margin-top: 2px; }
.stat .val.long { font-size: 11pt; font-weight: 650; letter-spacing: 0; line-height: 1.25; }
.stat .sub { font-size: 8pt; color: var(--faint); margin: 1px 0 0; }
.chart { margin: 8px 0 14px; page-break-inside: avoid; }
.chart .cap { font-size: 8.5pt; color: var(--faint); margin-top: 2px; }
.sources { font-size: 8.5pt; color: var(--muted); }
.sources li { margin-bottom: 2px; }
.footer { margin-top: 22px; border-top: 1px solid var(--hair); padding-top: 6px; display: flex; justify-content: space-between; gap: 12px; font-size: 8pt; color: var(--faint); }
.footer .conf { color: var(--alert); letter-spacing: 0; text-transform: uppercase; font-weight: 700; }
a { color: var(--accent-strong); text-decoration: none; }
.pos { color: var(--accent-strong); } .neg { color: var(--alert); }
"""


def html_document(title: str, body_html: str, kicker: str = "Tacite · Press Intelligence", subtitle: str = "", date_str: str | None = None) -> str:
    date_str = date_str or datetime.now().strftime("%d/%m/%Y")
    lines = [ln.strip() for ln in (subtitle or "").splitlines() if ln.strip()]
    if len(lines) >= 2:
        meta = "".join(f"<div>{_esc(ln)}</div>" for ln in lines)
        sub_html = f'<div class="meta">{meta}</div>'
        under = ""
    elif lines:
        sub_html = ""
        under = f'<p class="sub">{_esc(lines[0])}</p>'
    else:
        sub_html = under = ""
    css = GORGONE_CSS.replace("__FONT__", _font_css())
    return f"""<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>{_esc(title)}</title><style>{css}</style></head>
<body><div class="doc">
<header class="mast"><div class="mast-main"><div class="kicker">{_esc(kicker)}</div><h1>{_esc(title)}</h1>{under}</div>{sub_html}</header>
<hr class="rule">
{body_html}
<div class="footer"><span>Tacite — intelligence presse · {_esc(date_str)}</span><span class="conf">Usage interne</span></div>
</div></body></html>"""


def _esc(s) -> str:
    return (str(s) if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def pdf_from_html(html: str) -> bytes:
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.set_content(html, wait_until="load")
        pdf = pg.pdf(format="A4", print_background=True, margin={"top": "18mm", "bottom": "20mm", "left": "16mm", "right": "16mm"})
        b.close()
    return pdf


def docx_from_markdown(title: str, md: str, subtitle: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    k = doc.add_paragraph()
    r = k.add_run("TACITE · NOTE")
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x6E, 0x62, 0x30)
    doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].font.color.rgb = RGBColor(0x5E, 0x5A, 0x52)
    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        rows = [l for l in table_buf if not re.match(r"^\s*\|?\s*:?-{2,}", l)]
        cells = [[c.strip() for c in l.strip().strip("|").split("|")] for l in rows]
        if cells:
            t = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
            t.style = "Light List Accent 1" if "Light List Accent 1" in [s.name for s in doc.styles] else t.style
            for i, r in enumerate(cells):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
                    if i == 0:
                        for run in t.cell(i, j).paragraphs[0].runs:
                            run.font.bold = True
        table_buf = []

    for line in (md or "").splitlines():
        if line.strip().startswith("|"):
            table_buf.append(line)
            continue
        if table_buf:
            flush_table()
        s = line.rstrip()
        if not s:
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif re.match(r"^\s*[-*] ", s):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", s))
        elif re.match(r"^\s*\d+[.)] ", s):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)] ", "", s))
        elif s.startswith(">"):
            p = doc.add_paragraph()
            _add_runs(p, s.lstrip("> "))
            for run in p.runs:
                run.italic = True
        else:
            _add_runs(doc.add_paragraph(), s)
    if table_buf:
        flush_table()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _add_runs(p, text: str):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


# ───────────────────────────────────────────────────────────── graphiques SVG (pour les PDF)


def svg_radar(declared: dict, measured: dict, size: int = 380) -> str:
    cx = cy = size / 2
    r = size / 2 - 58
    n = len(AXES)
    def pt(i, v):
        ang = -math.pi / 2 + 2 * math.pi * i / n
        rr = r * (v + 1) / 2
        return cx + rr * math.cos(ang), cy + rr * math.sin(ang)
    grid = "".join(
        f'<polygon points="{" ".join(f"{cx + r * k * math.cos(-math.pi/2 + 2*math.pi*i/n):.1f},{cy + r * k * math.sin(-math.pi/2 + 2*math.pi*i/n):.1f}" for i in range(n))}" fill="none" stroke="#e3e0d8" stroke-width="1"/>'
        for k in (0.25, 0.5, 0.75, 1.0)
    )
    spokes = "".join(f'<line x1="{cx}" y1="{cy}" x2="{cx + r*math.cos(-math.pi/2 + 2*math.pi*i/n):.1f}" y2="{cy + r*math.sin(-math.pi/2 + 2*math.pi*i/n):.1f}" stroke="#e3e0d8"/>' for i in range(n))
    def poly(d, stroke, fill, dash=""):
        pts = " ".join(f"{x:.1f},{y:.1f}" for x, y in (pt(i, max(-1, min(1, float(d.get(a, 0) or 0)))) for i, a in enumerate(AXES)))
        return f'<polygon points="{pts}" fill="{fill}" stroke="{stroke}" stroke-width="1.5" {dash}/>'
    labels = ""
    for i, a in enumerate(AXES):
        ang = -math.pi / 2 + 2 * math.pi * i / n
        x, y = cx + (r + 22) * math.cos(ang), cy + (r + 22) * math.sin(ang)
        anchor = "middle" if abs(math.cos(ang)) < 0.2 else ("start" if math.cos(ang) > 0 else "end")
        labels += f'<text x="{x:.1f}" y="{y + 3:.1f}" font-size="9" fill="#5e5a52" text-anchor="{anchor}">{_esc(axis_labels()[a])}</text>'
    dashed = poly(declared, "#1a1916", "rgba(26,25,22,0.04)", 'stroke-dasharray="5 3"')
    solid = poly(measured, "#6e6230", "rgba(110,98,48,0.18)")
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{size}" height="{size}" viewBox="0 0 {size} {size}" font-family="Geist, Helvetica, Arial, sans-serif">'
            f'{grid}{spokes}{dashed}{solid}{labels}'
            f'<g font-size="9" fill="#5e5a52"><line x1="{cx-70}" y1="{size-10}" x2="{cx-50}" y2="{size-10}" stroke="#1a1916" stroke-dasharray="5 3"/><text x="{cx-46}" y="{size-7}">{"declared" if get_locale() == "en" else "déclaré"}</text>'
            f'<line x1="{cx+10}" y1="{size-10}" x2="{cx+30}" y2="{size-10}" stroke="#6e6230" stroke-width="2"/><text x="{cx+34}" y="{size-7}">{"measured" if get_locale() == "en" else "mesuré"}</text></g></svg>')


def svg_axis_bars(declared: dict, measured: dict, width: int = 620) -> str:
    """Une ligne par axe : pôle − à gauche, pôle + à droite, cercle creux = déclaré, point plein = mesuré."""
    row_h, left, right = 26, 150, 150
    h = len(AXES) * row_h + 30
    track = width - left - right
    x = lambda v: left + (max(-1, min(1, v)) + 1) / 2 * track
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{h}" viewBox="0 0 {width} {h}" font-family="Geist, Helvetica, Arial, sans-serif">']
    out.append(f'<g font-size="8" fill="#9a958b" text-anchor="middle">' + "".join(f'<text x="{x(v):.1f}" y="10">{"+" if v > 0 else ""}{v:.1f}</text>' for v in (-1, -0.5, 0, 0.5, 1)) + "</g>")
    for i, a in enumerate(AXES):
        cy = 20 + i * row_h + row_h / 2
        neg, pos = axis_poles_map()[a]
        d, m = float(declared.get(a, 0) or 0), float(measured.get(a, 0) or 0)
        if i % 2 == 0:
            out.append(f'<rect x="0" y="{cy - row_h/2:.1f}" width="{width}" height="{row_h}" fill="#f3f1ec" opacity="0.6"/>')
        out.append(f'<text x="{left-8}" y="{cy-3:.1f}" font-size="8.5" font-weight="700" fill="#6e6230" text-anchor="end" letter-spacing="0">{_esc(axis_labels()[a].upper())}</text>')
        out.append(f'<text x="{left-8}" y="{cy+8:.1f}" font-size="8" fill="#9a958b" text-anchor="end">◂ {_esc(neg)}</text>')
        out.append(f'<text x="{width-right+8}" y="{cy+3:.1f}" font-size="8" fill="#9a958b">{_esc(pos)} ▸</text>')
        out.append(f'<line x1="{left}" x2="{width-right}" y1="{cy:.1f}" y2="{cy:.1f}" stroke="#e3e0d8"/>')
        out.append(f'<line x1="{x(0):.1f}" x2="{x(0):.1f}" y1="{cy-8:.1f}" y2="{cy+8:.1f}" stroke="#c9c4b8" stroke-dasharray="2 2"/>')
        out.append(f'<line x1="{x(d):.1f}" x2="{x(m):.1f}" y1="{cy:.1f}" y2="{cy:.1f}" stroke="#1a1916" opacity="0.5"/>')
        out.append(f'<circle cx="{x(d):.1f}" cy="{cy:.1f}" r="4" fill="#fff" stroke="#1a1916" stroke-dasharray="1.5 1.5"/>')
        out.append(f'<circle cx="{x(m):.1f}" cy="{cy:.1f}" r="5" fill="#6e6230"/>')
    out.append("</svg>")
    return "".join(out)


def svg_hbars(items: list[tuple[str, float]], max_value: float | None = None, width: int = 620, fmt: str = "{:.0f}") -> str:
    row_h, left = 22, 190
    mx = max_value or max([abs(v) for _, v in items] + [1])
    h = len(items) * row_h + 8
    track = width - left - 60
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{h}" viewBox="0 0 {width} {h}" font-family="Geist, Helvetica, Arial, sans-serif">']
    for i, (lab, v) in enumerate(items):
        cy = 4 + i * row_h + row_h / 2
        w = abs(v) / mx * track
        out.append(f'<text x="{left-8}" y="{cy+3.5:.1f}" font-size="9.5" fill="#1a1916" text-anchor="end">{_esc(lab[:34])}</text>')
        out.append(f'<rect x="{left}" y="{cy-6:.1f}" width="{track}" height="12" fill="#f3f1ec"/>')
        out.append(f'<rect x="{left}" y="{cy-6:.1f}" width="{w:.1f}" height="12" fill="{"#8a2b1d" if v < 0 else "#6e6230"}"/>')
        out.append(f'<text x="{left+track+8}" y="{cy+3.5:.1f}" font-size="9" fill="#5e5a52">{fmt.format(v)}</text>')
    out.append("</svg>")
    return "".join(out)
